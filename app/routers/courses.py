import csv
import io
import uuid
from datetime import datetime
from typing import Literal, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.dependencies import get_current_user_id, require_professor
from app.database import get_db
from app.models.course import Course, Enrollment
from app.models.session import Session
from app.models.user import User
from app.schemas.course import (
    CourseListResponse,
    CourseOut,
    CourseStatsResponse,
    CreateCourseRequest,
    UpdateCourseRequest,
)
from app.schemas.session import SessionOut
from app.services import course as course_service

router = APIRouter()


class EnrollRequest(BaseModel):
    identifier: str  # email or matric number


async def _get_owned_course(course_id: str, professor_id: str, db: AsyncSession) -> Course:
    course = await course_service.get_course_by_id(db, course_id)
    if not course:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Course not found")
    if str(course.professor_id) != professor_id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Not your course")
    return course


@router.get("", response_model=CourseListResponse)
async def list_courses(
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    user = await db.get(User, uuid.UUID(user_id))
    if not user:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="User not found")

    if user.role == "professor":
        courses = await course_service.get_courses_for_professor(db, user_id)
    else:
        courses = await course_service.get_courses_for_student(db, user_id)

    return CourseListResponse(courses=courses)


@router.post("", response_model=CourseOut, status_code=status.HTTP_201_CREATED)
async def create_course(
    body: CreateCourseRequest,
    professor_id: str = Depends(require_professor),
    db: AsyncSession = Depends(get_db),
):
    course = await course_service.create_course(db, professor_id, body)
    return CourseOut(id=course.id, code=course.code, name=course.name, studentCount=0)


@router.get("/{course_id}", response_model=CourseOut)
async def get_course(
    course_id: str,
    professor_id: str = Depends(require_professor),
    db: AsyncSession = Depends(get_db),
):
    course = await _get_owned_course(course_id, professor_id, db)
    cnt = (await db.execute(
        select(func.count(Enrollment.student_id)).where(Enrollment.course_id == course.id)
    )).scalar_one() or 0
    return CourseOut(id=course.id, code=course.code, name=course.name, studentCount=cnt)


@router.put("/{course_id}", response_model=CourseOut)
async def update_course(
    course_id: str,
    body: UpdateCourseRequest,
    professor_id: str = Depends(require_professor),
    db: AsyncSession = Depends(get_db),
):
    course = await _get_owned_course(course_id, professor_id, db)
    course = await course_service.update_course(db, course, body)
    cnt = (await db.execute(
        select(func.count(Enrollment.student_id)).where(Enrollment.course_id == course.id)
    )).scalar_one() or 0
    return CourseOut(id=course.id, code=course.code, name=course.name, studentCount=cnt)


@router.delete("/{course_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_course(
    course_id: str,
    professor_id: str = Depends(require_professor),
    db: AsyncSession = Depends(get_db),
):
    course = await _get_owned_course(course_id, professor_id, db)
    await course_service.delete_course(db, course)


@router.get("/{course_id}/sessions", response_model=list[SessionOut])
async def list_course_sessions(
    course_id: str,
    professor_id: str = Depends(require_professor),
    db: AsyncSession = Depends(get_db),
):
    course = await _get_owned_course(course_id, professor_id, db)
    rows = (await db.execute(
        select(Session)
        .where(Session.course_id == course.id)
        .order_by(Session.started_at.desc())
    )).scalars().all()
    return [
        SessionOut(
            id=s.id,
            courseId=s.course_id,
            status=s.status,
            startedAt=s.started_at,
            endedAt=s.ended_at,
        )
        for s in rows
    ]


@router.post("/{course_id}/enroll", status_code=status.HTTP_201_CREATED)
async def enroll_student(
    course_id: str,
    body: EnrollRequest,
    professor_id: str = Depends(require_professor),
    db: AsyncSession = Depends(get_db),
):
    await _get_owned_course(course_id, professor_id, db)
    student = await course_service.enroll_student(db, course_id, body.identifier)
    return {"student_id": str(student.id), "student_name": student.full_name, "email": student.email}


@router.delete("/{course_id}/students/{student_id}", status_code=status.HTTP_204_NO_CONTENT)
async def remove_student(
    course_id: str,
    student_id: str,
    professor_id: str = Depends(require_professor),
    db: AsyncSession = Depends(get_db),
):
    await _get_owned_course(course_id, professor_id, db)
    await course_service.remove_student(db, course_id, student_id)


@router.get("/{course_id}/stats", response_model=CourseStatsResponse)
async def get_course_stats(
    course_id: str,
    professor_id: str = Depends(require_professor),
    db: AsyncSession = Depends(get_db),
):
    course = await _get_owned_course(course_id, professor_id, db)
    total_sessions, students = await course_service.get_course_stats(db, str(course.id))
    return CourseStatsResponse(
        course_id=course.id,
        total_sessions=total_sessions,
        students=students,
    )


@router.get("/{course_id}/attendance/export")
async def export_attendance(
    course_id: str,
    professor_id: str = Depends(require_professor),
    db: AsyncSession = Depends(get_db),
    format: Literal["csv", "xlsx"] = Query("csv"),
    from_date: Optional[str] = Query(None),
    to_date: Optional[str] = Query(None),
):
    course = await _get_owned_course(course_id, professor_id, db)
    records = await course_service.export_course_attendance(db, course_id, from_date, to_date)
    filename = f"attendance_{course.code}"

    exported_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    date_range = f"{from_date or 'all'} to {to_date or 'all'}"

    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([f"Course: {course.code} — {course.name}"])
        writer.writerow([f"Exported: {exported_at}"])
        writer.writerow([f"Date Range: {date_range}"])
        writer.writerow([f"Total Records: {len(records)}"])
        writer.writerow([])
        writer.writerow(["Student Name", "Matric No", "Session Date", "Marked At"])
        for r in records:
            writer.writerow([
                r["student_name"],
                r["matric_no"],
                r["session_date"].isoformat() if r["session_date"] else "",
                r["marked_at"].isoformat() if r["marked_at"] else "",
            ])
        output.seek(0)
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{filename}.csv"'},
        )

    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    meta_rows = [
        [f"Course: {course.code} — {course.name}"],
        [f"Exported: {exported_at}"],
        [f"Date Range: {date_range}"],
        [f"Total Records: {len(records)}"],
        [],
    ]
    for row in meta_rows:
        ws.append(row)
    header_row = ws.max_row + 1
    ws.append(["Student Name", "Matric No", "Session Date", "Marked At"])
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
    for r in records:
        ws.append([
            r["student_name"],
            r["matric_no"],
            r["session_date"].isoformat() if r["session_date"] else "",
            r["marked_at"].isoformat() if r["marked_at"] else "",
        ])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"'},
    )
